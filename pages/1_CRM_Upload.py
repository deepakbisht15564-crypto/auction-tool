from datetime import datetime
import os
import subprocess
import time
import zipfile
import openpyxl
from playwright.sync_api import sync_playwright
import streamlit as st

# Chromium install check
try:
  subprocess.run(['playwright', 'install', 'chromium'], check=True)
except Exception:
  pass

st.set_page_config(page_title='CRM Auto Uploader', layout='wide')
st.title('🚀 BFSI CRM Auto Uploader (With OTP Support)')
st.caption(
    'Aapka main PDF extractor bilkul safe hai. Ye module extracted ZIP ko CRM me'
    ' upload karega.'
)

CRM_USER = st.secrets.get('CRM_USER', '')
CRM_PASS = st.secrets.get('CRM_PASS', '')
LOGIN_URL = 'https://crm.auctionbaba.com/login'
FORM_URL = 'https://crm.auctionbaba.com/Xarc/new_account_creation_reff'

# Date Calculations
today = datetime.today()
TODAY_DATE = today.strftime('%Y-%B-%d')
fy_year = today.year + 1 if today.month > 3 else today.year
FY_END_DATE = f'{fy_year}-March-31'

# Session State for bot persistent state
if 'browser_session' not in st.session_state:
  st.session_state.browser_session = None
if 'otp_sent' not in st.session_state:
  st.session_state.otp_sent = False

# 1. Upload Extracted Deliverable ZIP
uploaded_zip = st.file_uploader(
    'Upload Generated ZIP File (Contains Styled Excel + Notice Cuttings)',
    type=['zip'],
)

if uploaded_zip:
  # Unzip to temporary folder
  temp_dir = '/tmp/crm_batch'
  os.makedirs(temp_dir, exist_ok=True)
  with zipfile.ZipFile(uploaded_zip, 'r') as zip_ref:
    zip_ref.extractall(temp_dir)

  # Find Excel file inside ZIP
  excel_files = [f for f in os.listdir(temp_dir) if f.endswith('.xlsx')]
  if not excel_files:
    st.error('ZIP file me koi .xlsx file nahi mili!')
  else:
    excel_path = os.path.join(temp_dir, excel_files[0])
    wb = openpyxl.load_workbook(excel_path)
    bank_sheets = [s for s in wb.sheetnames if s.strip() != 'Master']

    st.success(
        f'Excel Loaded! Found {len(bank_sheets)} Bank Sheets:'
        f" {', '.join(bank_sheets)}"
    )

    st.markdown('---')
    st.subheader('Step 1: CRM Login & Trigger OTP')

    col1, col2 = st.columns([2, 1])

    with col1:
      if st.button(
          '📲 Send OTP to My Mobile', disabled=st.session_state.otp_sent
      ):
        with st.spinner('Connecting to CRM and requesting OTP...'):
          p = sync_playwright().start()
          browser = p.chromium.launch(
              headless=True,
              args=['--no-sandbox', '--disable-dev-shm-usage', '--disable-gpu'],
          )
          context = browser.new_context()
          page = context.new_page()

          page.goto(LOGIN_URL)
          page.wait_for_timeout(2000)

          # Username & Password Entry
          page.fill("input[name*='user'], input[id*='user']", CRM_USER)
          page.fill("input[name*='pass'], input[id*='pass']", CRM_PASS)
          page.click("button[type='submit'], input[type='submit']")
          page.wait_for_timeout(3000)

          # Store objects in session
          st.session_state['playwright_p'] = p
          st.session_state['playwright_browser'] = browser
          st.session_state['playwright_page'] = page
          st.session_state.otp_sent = True
          st.success(
              'OTP aapke mobile par bhej diya gaya hai! Niche enter karein.'
          )
          st.rerun()

    # Step 2: OTP Entry & Push to CRM
    if st.session_state.otp_sent:
      st.markdown('---')
      st.subheader('Step 2: Enter Mobile OTP & Start Upload')

      otp_val = st.text_input(
          'Enter Mobile OTP:',
          max_chars=6,
          placeholder='e.g. 123456',
          type='password',
      )

      if st.button('🚀 Verify OTP & Push All Accounts'):
        if not otp_val:
          st.warning('Pehle OTP enter kijiye!')
        else:
          page = st.session_state.get('playwright_page')
          browser = st.session_state.get('playwright_browser')
          p = st.session_state.get('playwright_p')

          with st.spinner('Logging in and creating accounts in CRM...'):
            try:
              # OTP fill & Submit
              page.fill(
                  "input[name*='otp'], input[id*='otp'], input[type='number']",
                  otp_val,
              )
              page.click(
                  "button:has-text('Verify'), button:has-text('Submit'),"
                  " input[type='submit']"
              )
              page.wait_for_timeout(4000)

              # Navigate to New Account Creation Page
              page.goto(FORM_URL)
              page.wait_for_selector(
                  'text=Client Account Assignment', timeout=40000
              )

              st.info('Login Successful! Starting Bank Batches...')

              progress_bar = st.progress(0)
              status_box = st.empty()

              for b_idx, sheet_name in enumerate(bank_sheets):
                sheet = wb[sheet_name]
                rows = list(sheet.iter_rows(values_only=True))
                if len(rows) <= 1:
                  continue

                headers = [str(h).strip() if h else '' for h in rows[0]]
                data_rows = rows[1:]
                total_accounts = len(data_rows)
                bank_name = sheet_name.strip()

                status_box.text(
                    f'Processing [{bank_name}] ({total_accounts} Accounts)...'
                )

                if page.url != FORM_URL:
                  page.goto(FORM_URL)
                  page.wait_for_selector('text=Client Account Assignment')

                # Top Dropdowns
                page.locator('select').nth(0).select_option(
                    label='National Head - KunDan Kumar Jha ( bd )'
                )
                try:
                  page.locator('select').nth(1).select_option(label=bank_name)
                except Exception:
                  page.locator('select').nth(1).select_option(value=bank_name)

                page.locator('select').nth(2).select_option(label='Secured')
                page.locator('select').nth(3).select_option(
                    label='Bidder Sourcing'
                )

                # Notice Cutting Upload
                cutting_files = [
                    os.path.join(temp_dir, f)
                    for f in os.listdir(temp_dir)
                    if bank_name.lower() in f.lower()
                    and f.endswith(('.png', '.jpg', '.jpeg'))
                ]
                if cutting_files:
                  page.locator("input[type='file']").first.set_input_files(
                      cutting_files[0]
                  )

                # Dates
                page.locator("input[name*='start_date']").first.fill(TODAY_DATE)
                page.locator("input[name*='end_date']").first.fill(FY_END_DATE)

                # Add Accounts Rows
                page.locator(
                    "input[name*='add_account'], #add_account"
                ).first.fill(str(total_accounts))
                page.locator(
                    "button:has-text('ADD'), input[value='ADD']"
                ).click()
                page.wait_for_timeout(2500)

                # Fill Borrower Rows (Skipping Box 1 and Box 2)
                col_map = {
                    name.lower(): idx for idx, name in enumerate(headers)
                }
                b_idx_col = col_map.get('borrower name', 2)
                d_idx_col = col_map.get('property description', 3)
                p_idx_col = col_map.get('reserve price', 4)

                for r_idx, r_data in enumerate(data_rows, start=1):
                  b_name = str(r_data[b_idx_col] or '').strip()
                  p_desc = str(r_data[d_idx_col] or '').strip()
                  r_price = str(r_data[p_idx_col] or '').strip()

                  # Account Name
                  page.locator(
                      f'tr:nth-child({r_idx}) input[name*="account"]'
                  ).fill(b_name)
                  # Asset Description
                  page.locator(f'tr:nth-child({r_idx}) textarea').fill(p_desc)
                  # Reserve Price
                  page.locator(
                      f'tr:nth-child({r_idx}) input[name*="reserve"]'
                  ).fill(r_price)

                # Submit Batch
                page.locator(
                    "button:has-text('Submit'), input[value='Submit']"
                ).click()
                page.wait_for_timeout(3500)

                st.write(
                    f'✅ **{bank_name}**: {total_accounts} accounts'
                    ' successfully created!'
                )
                progress_bar.progress((b_idx + 1) / len(bank_sheets))

              st.balloons()
              st.success('Sabhi banks ka data successfully CRM me chala gaya!')

            except Exception as err:
              st.error(f'Error during upload: {str(err)}')
            finally:
              browser.close()
              p.stop()
              st.session_state.otp_sent = False
